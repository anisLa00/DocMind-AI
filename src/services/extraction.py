"""Turn an uploaded file into plain text, one entry per page."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from src.utils.errors import UnsupportedFileTypeError, ValidationError

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".csv"}


@dataclass(frozen=True)
class ExtractedPage:
    """Text of a single page. ``number`` is 1-indexed; 1 for page-less formats."""

    number: int
    text: str


def _clean(text: str) -> str:
    """Normalise line endings and collapse the runs of blank lines PDFs emit."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [line.rstrip() for line in text.split("\n")]

    cleaned: list[str] = []
    blank_run = 0
    for line in lines:
        if line:
            blank_run = 0
            cleaned.append(line)
            continue
        blank_run += 1
        if blank_run <= 1:
            cleaned.append("")

    return "\n".join(cleaned).strip()


def _extract_pdf(path: Path) -> list[ExtractedPage]:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(str(path))
    except PyPdfError as exc:
        raise ValidationError(f"Could not read PDF: {exc}") from exc

    if reader.is_encrypted:
        # An empty user password is common for "protected" PDFs and is worth a try.
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValidationError("PDF is password protected") from exc

    pages: list[ExtractedPage] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = _clean(page.extract_text() or "")
        except Exception:
            text = ""
        if text:
            pages.append(ExtractedPage(number=index, text=text))

    return pages


def _extract_docx(path: Path) -> list[ExtractedPage]:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(str(path))
    except PackageNotFoundError as exc:
        raise ValidationError("Could not read DOCX: file is not a valid Word document") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = _clean("\n".join(parts))
    # python-docx exposes no page breaks, so a DOCX is treated as a single page.
    return [ExtractedPage(number=1, text=text)] if text else []


def _extract_text_file(path: Path) -> list[ExtractedPage]:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = _clean(raw.decode(encoding))
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes any byte string
        raise ValidationError("Could not decode text file")

    return [ExtractedPage(number=1, text=text)] if text else []


def extract_pages(file_path: str | Path) -> list[ExtractedPage]:
    """Extract text from a supported document, skipping pages with no text."""
    path = Path(file_path)
    if not path.is_file():
        raise ValidationError(f"File not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in TEXT_EXTENSIONS:
        return _extract_text_file(path)

    raise UnsupportedFileTypeError(f"Unsupported file type: {suffix or path.name}")


def extract_text(file_path: str | Path) -> str:
    """Whole-document text, pages joined by blank lines."""
    return "\n\n".join(page.text for page in extract_pages(file_path))
