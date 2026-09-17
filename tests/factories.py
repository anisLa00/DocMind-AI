"""Builders for the binary document formats the extractor supports."""

import io


def build_pdf(pages: list[list[str]]) -> bytes:
    """A minimal, uncompressed multi-page PDF with selectable text."""
    objects: list[bytes] = []
    page_object_ids: list[int] = []

    # 1 = catalog, 2 = page tree, 3 = font; pages and their content follow.
    next_id = 4
    for lines in pages:
        content = "BT /F1 12 Tf 72 720 Td 14 TL\n"
        for line in lines:
            escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            content += f"({escaped}) Tj T*\n"
        content += "ET"
        stream = content.encode("latin-1")

        page_id, content_id = next_id, next_id + 1
        next_id += 2
        page_object_ids.append(page_id)

        objects.append(
            (
                page_id,
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 3 0 R >> >> /Contents "
                + f"{content_id} 0 R >>".encode(),
            )
        )
        objects.append(
            (
                content_id,
                b"<< /Length "
                + str(len(stream)).encode()
                + b" >>\nstream\n"
                + stream
                + b"\nendstream",
            )
        )

    kids = " ".join(f"{page_id} 0 R" for page_id in page_object_ids)
    header = [
        (1, b"<< /Type /Catalog /Pages 2 0 R >>"),
        (
            2,
            f"<< /Type /Pages /Kids [{kids}] /Count {len(page_object_ids)} >>".encode(),
        ),
        (3, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"),
    ]

    all_objects = sorted(header + objects)
    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}

    for object_id, body in all_objects:
        offsets[object_id] = len(out)
        out += f"{object_id} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_at = len(out)
    count = len(all_objects) + 1
    out += f"xref\n0 {count}\n".encode() + b"0000000000 65535 f \n"
    for object_id in range(1, count):
        out += f"{offsets[object_id]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()

    return bytes(out)


def build_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    if table:
        word_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                word_table.cell(row_index, column_index).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
